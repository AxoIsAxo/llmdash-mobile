from __future__ import annotations
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

from .ast_types import (
    BlockNode, InlineNode, Text, Bold, Italic, Strikethrough, Underline,
    InlineCode, Link, Image, MathInline, Subscript, Superscript,
    Heading, Paragraph, BulletList, ListItem, Table as TableNode, CodeBlock,
    BlockQuote, HorizontalRule, ChartDirective, MathBlock,
)
from .chart_generator import generate_chart_png


def build_docx(
    nodes: list[BlockNode],
    filepath: str,
    font_family: str = "serif",
    heading_font: str = "sans",
    font_size: int = 11,
    charts_dir: str = "data/documents/charts",
):
    doc = Document()

    _setup_default_styles(doc, font_family, heading_font, font_size)

    for node in nodes:
        _add_block(doc, node, font_family, heading_font, font_size, charts_dir)

    doc.save(filepath)


def _setup_default_styles(doc: Document, font_family: str, heading_font: str, font_size: int):
    style = doc.styles["Normal"]
    style.font.size = Pt(font_size)
    _set_font(style, font_family)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    for level in range(1, 7):
        h_style = doc.styles[f"Heading {level}"]
        _set_font(h_style, heading_font)
        h_style.font.size = Pt(max(24 - (level * 2), 12))
        h_style.font.bold = True
        h_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        h_style.paragraph_format.space_before = Pt(18 - (level * 2))
        h_style.paragraph_format.space_after = Pt(8 - level)

    list_style = doc.styles["List Bullet"]
    _set_font(list_style, font_family)
    list_style.font.size = Pt(font_size)

    list_num_style = doc.styles["List Number"]
    _set_font(list_num_style, font_family)
    list_num_style.font.size = Pt(font_size)


def _set_font(style, family: str):
    font_map = {
        "serif": "Georgia",
        "sans": "Calibri",
        "mono": "Consolas",
    }
    name = font_map.get(family, family)
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:eastAsia="{name}" w:cs="{name}"/>')
        rpr.append(rFonts)
    else:
        rFonts.set(qn("w:ascii"), name)
        rFonts.set(qn("w:hAnsi"), name)


def _add_block(doc, node: BlockNode, font_family: str, heading_font: str, font_size: int, charts_dir: str):
    if isinstance(node, Heading):
        p = doc.add_heading("", level=node.level)
        _add_inline_runs(p, node.children, font_family, heading_font)

    elif isinstance(node, Paragraph):
        p = doc.add_paragraph()
        _add_inline_runs(p, node.children, font_family, heading_font)

    elif isinstance(node, BulletList):
        _add_list(doc, node, font_family, heading_font, font_size, charts_dir, level=0)

    elif isinstance(node, TableNode):
        _add_table(doc, node, font_family, heading_font)

    elif isinstance(node, CodeBlock):
        _add_code_block(doc, node, font_family, font_size)

    elif isinstance(node, BlockQuote):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.right_indent = Cm(0.5)
        run = p.add_run()
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        inner_text = _inline_children_to_text(node.children)
        run.text = inner_text

    elif isinstance(node, HorizontalRule):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

    elif isinstance(node, ChartDirective):
        chart_path = generate_chart_png(node, charts_dir)
        if chart_path and os.path.exists(chart_path):
            doc.add_picture(chart_path, width=Inches(5))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    elif isinstance(node, MathBlock):
        p = doc.add_paragraph()
        run = p.add_run(node.latex)
        run.font.italic = True
        run.font.size = Pt(font_size + 1)


def _add_inline_runs(p, children: list[InlineNode], font_family: str, heading_font: str):
    for child in children:
        _append_run(p, child, font_family, heading_font)


def _append_run(p, node: InlineNode, font_family: str, heading_font: str):
    if isinstance(node, Text):
        run = p.add_run(node.text)
    elif isinstance(node, Bold):
        run = p.add_run(_inline_children_to_text(node.children))
        run.font.bold = True
    elif isinstance(node, Italic):
        run = p.add_run(_inline_children_to_text(node.children))
        run.font.italic = True
    elif isinstance(node, Strikethrough):
        run = p.add_run(_inline_children_to_text(node.children))
        run.font.strike = True
    elif isinstance(node, Underline):
        run = p.add_run(_inline_children_to_text(node.children))
        run.font.underline = True
    elif isinstance(node, InlineCode):
        run = p.add_run(node.text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F0F0F0"/>')
        run._r.get_or_add_rPr().append(shd)
    elif isinstance(node, Link):
        run = p.add_run(node.text)
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.font.underline = True
    elif isinstance(node, Image):
        if os.path.exists(node.url):
            p2 = p
            try:
                p2 = p.insert_paragraph_before()
            except Exception:
                p2 = p
            try:
                p2.add_run().add_picture(node.url, width=Inches(4))
            except Exception:
                pass
    elif isinstance(node, MathInline):
        run = p.add_run(f"${node.latex}$")
        run.font.italic = True


def _add_list(doc, node: BulletList, font_family: str, heading_font: str, font_size: int, charts_dir: str, level: int):
    for item in node.items:
        if node.ordered:
            p = doc.add_paragraph(style="List Number")
        else:
            p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(1.0 + level * 0.8)
        inner_text = _list_item_text(item, font_family, heading_font)
        run = p.add_run(inner_text)
        run.font.size = Pt(font_size)
        _add_nested_lists(doc, item, font_family, heading_font, font_size, charts_dir, level + 1)


def _add_nested_lists(doc, item: ListItem, font_family: str, heading_font: str, font_size: int, charts_dir: str, level: int):
    for child in item.children:
        if isinstance(child, BulletList):
            _add_list(doc, child, font_family, heading_font, font_size, charts_dir, level)


def _add_table(doc, node: TableNode, font_family: str, heading_font: str):
    num_cols = max(len(node.headers) if node.headers else 0, max((len(r) for r in node.rows), default=0))
    if num_cols == 0:
        return
    table = doc.add_table(rows=1 + len(node.rows), cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if node.headers:
        for i, cell_content in enumerate(node.headers):
            cell = table.rows[0].cells[i]
            cell._tc.get_or_add_tcPr()
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8E8E8" w:val="clear"/>')
            cell._tc.get_or_add_tcPr().append(shading)
            p = cell.paragraphs[0]
            run = p.add_run(_inline_children_to_text(cell_content))
            run.font.bold = True
            run.font.size = Pt(10)

    for ri, row in enumerate(node.rows):
        for ci, cell_content in enumerate(row):
            if ci >= num_cols:
                continue
            cell = table.rows[ri + 1].cells[ci]
            p = cell.paragraphs[0]
            run = p.add_run(_inline_children_to_text(cell_content))
            run.font.size = Pt(10)


def _add_code_block(doc, node: CodeBlock, font_family: str, font_size: int):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5" w:val="clear"/>')
    pPr.append(shd)
    run = p.add_run(node.code)
    run.font.name = "Consolas"
    run.font.size = Pt(font_size - 2)


def _inline_children_to_text(children: list[InlineNode]) -> str:
    parts: list[str] = []
    for c in children:
        if isinstance(c, Text):
            parts.append(c.text)
        elif isinstance(c, InlineCode):
            parts.append(c.text)
        elif isinstance(c, (Bold, Italic, Strikethrough, Underline)):
            parts.append(_inline_children_to_text(c.children))
        elif isinstance(c, Link):
            parts.append(c.text)
        elif isinstance(c, MathInline):
            parts.append(f"${c.latex}$")
        elif isinstance(c, (Subscript, Superscript)):
            parts.append(_inline_children_to_text(c.children))
        elif isinstance(c, Image):
            parts.append(c.alt or "")
    return "".join(parts)


def _list_item_text(item: ListItem, font_family: str, heading_font: str) -> str:
    parts: list[str] = []
    for child in item.children:
        if isinstance(child, Paragraph):
            parts.append(_inline_children_to_text(child.children))
        elif isinstance(child, BulletList):
            pass
        elif isinstance(child, CodeBlock):
            parts.append(child.code)
    return "".join(parts)
